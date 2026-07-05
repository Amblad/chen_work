"""
从分段式考勤 Excel 报表中提取数据行，输出到新的 Excel 文件。

用法:
  python extract_attendance.py <输入文件或文件夹> <输出文件.xlsx>

示例:
  python extract_attendance.py report.xlsx output.xlsx
  python extract_attendance.py ./reports ./output.xlsx
"""

import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.read_only import EmptyCell

DATE_PATTERN = re.compile(r"^\d{1,2}[-/]\d{1,2}[-/]\d{4}$")
NAME_PATTERN = re.compile(r"Name\s*:\s*(.+)", re.IGNORECASE)
NAME_LABEL_PATTERN = re.compile(r"^Name\s*:?\s*$", re.IGNORECASE)

# 原文件列号 (1-based)
COL_DATE = 1       # A
COL_SCHED_IN = 3   # C
COL_SCHED_OUT = 4  # D
COL_BREAK = 5      # E
COL_TIME_IN = 7    # G
COL_TIME_OUT = 8   # H
COL_BRANCH = 27    # AA
COL_REASON = 28    # AB


def is_date_value(value) -> bool:
    if value is None:
        return False
    if isinstance(value, (datetime, date)):
        return True
    if isinstance(value, (int, float)):
        # Excel 日期序列值，大致对应 1982-2064 年
        n = float(value)
        return 30000 <= n <= 60000
    text = str(value).strip()
    if not text or text.lower().startswith("total"):
        return False
    if DATE_PATTERN.match(text):
        return True
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d-%B-%Y"):
        try:
            datetime.strptime(text, fmt)
            return True
        except ValueError:
            continue
    return False


def extract_name_from_row(cells) -> str | None:
    """从员工信息行中提取 Name 值。"""
    values = []
    for cell in cells:
        if isinstance(cell, EmptyCell):
            values.append(None)
        else:
            values.append(cell.value)

    for i, value in enumerate(values):
        if value is None:
            continue
        text = str(value).strip()

        match = NAME_PATTERN.search(text)
        if match:
            name = match.group(1).strip()
            if name:
                return name

        if NAME_LABEL_PATTERN.match(text) and i + 1 < len(values) and values[i + 1]:
            return str(values[i + 1]).strip()

    return None


def row_column_map(row_cells) -> dict[int, object]:
    return {
        cell.column: cell.value
        for cell in row_cells
        if not isinstance(cell, EmptyCell)
    }


def process_workbook(filepath: Path) -> list[list]:
    suffix = filepath.suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError(
            f"不支持的文件格式: {filepath.name}，请使用 .xlsx 或 .xlsm 文件"
        )

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"无法打开文件 {filepath.name}: {exc}") from exc
    sheet = wb.active
    filename = filepath.stem
    results = []
    current_name = None

    for row_cells in sheet.iter_rows():
        col = row_column_map(row_cells)

        name = extract_name_from_row(row_cells)
        if name:
            current_name = name
            continue

        date_val = col.get(COL_DATE)
        if not is_date_value(date_val):
            continue

        time_in = col.get(COL_TIME_IN)
        time_out = col.get(COL_TIME_OUT)

        results.append(
            [
                filename,
                col.get(COL_BRANCH),
                date_val,
                current_name,
                None,
                None,
                col.get(COL_SCHED_IN),
                col.get(COL_SCHED_OUT),
                col.get(COL_BREAK),
                time_in,
                time_out,
                col.get(COL_REASON),
                None,
                time_in,
                time_out,
            ]
        )

    wb.close()
    return results


def collect_input_files(path: Path) -> list[Path]:
    if path.is_file():
        return [path]

    files = []
    for pattern in ("*.xlsx", "*.xlsm"):
        files.extend(path.glob(pattern))
    return sorted(f for f in files if not f.name.startswith("~$"))


def write_output(rows: list[list], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Extracted"

    for row in rows:
        ws.append(row)

    try:
        wb.save(output_path)
    except PermissionError as exc:
        raise PermissionError(
            f"无法保存到 {output_path}，请先关闭正在打开的 Excel 文件后重试"
        ) from exc


def run_extraction(
    input_path: Path,
    output_path: Path,
    log=print,
) -> int:
    """执行提取，返回总行数。"""
    if not input_path.exists():
        raise FileNotFoundError(f"输入路径不存在: {input_path}")

    input_files = collect_input_files(input_path)
    if not input_files:
        raise FileNotFoundError(f"未找到 Excel 文件: {input_path}")

    all_rows = []
    for filepath in input_files:
        log(f"处理: {filepath.name}")
        rows = process_workbook(filepath)
        log(f"  提取 {len(rows)} 行")
        all_rows.extend(rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_output(all_rows, output_path)
    log(f"完成: 共 {len(all_rows)} 行 -> {output_path}")
    return len(all_rows)


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 1

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    try:
        run_extraction(input_path, output_path)
    except FileNotFoundError as exc:
        print(f"错误: {exc}")
        return 1

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
